import os
import json
import sqlite3
from datetime import datetime
from flask import Flask, render_template, request, jsonify, redirect, url_for, session, send_file, send_from_directory
from flask_login import LoginManager, UserMixin, login_user, logout_user, login_required, current_user
from werkzeug.security import generate_password_hash, check_password_hash
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import io

app = Flask(__name__)
app.secret_key = 'change-this-secret-key-in-production-ge-report-2026'

# Simple admin user (change password!)
ADMIN_USERNAME = 'admin'
ADMIN_PASSWORD_HASH = generate_password_hash('admin123')  # Change this!

login_manager = LoginManager()
login_manager.init_app(app)
login_manager.login_view = 'login'

class User(UserMixin):
    def __init__(self, id):
        self.id = id

@login_manager.user_loader
def load_user(user_id):
    if user_id == '1':
        return User('1')
    return None

# Database
DB_PATH = 'ge_reports.db'
UPLOAD_FOLDER = 'uploads'
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER

def init_db():
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    c.execute('''
        CREATE TABLE IF NOT EXISTS reports (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            submitted_at TEXT,
            batiment TEXT,
            redacteur TEXT,
            date_essai TEXT,
            data_json TEXT,
            photo1 TEXT,
            photo2 TEXT
        )
    ''')
    conn.commit()
    conn.close()

init_db()

def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn

@app.route('/')
def index():
    return render_template('form.html')

@app.route('/submit', methods=['POST'])
def submit():
    try:
        # Get JSON data
        data = request.form.get('data')
        if not data:
            return jsonify({'success': False, 'error': 'No data'}), 400
        data_dict = json.loads(data)
        
        # Handle photos
        photo1_filename = None
        photo2_filename = None
        if 'photo1' in request.files:
            photo1 = request.files['photo1']
            if photo1 and photo1.filename:
                photo1_filename = f"{datetime.now().strftime('%Y%m%d_%H%M%S')}_photo1_{photo1.filename}"
                photo1.save(os.path.join(app.config['UPLOAD_FOLDER'], photo1_filename))
        if 'photo2' in request.files:
            photo2 = request.files['photo2']
            if photo2 and photo2.filename:
                photo2_filename = f"{datetime.now().strftime('%Y%m%d_%H%M%S')}_photo2_{photo2.filename}"
                photo2.save(os.path.join(app.config['UPLOAD_FOLDER'], photo2_filename))
        
        conn = get_db()
        c = conn.cursor()
        c.execute('''
            INSERT INTO reports (submitted_at, batiment, redacteur, date_essai, data_json, photo1, photo2)
            VALUES (?, ?, ?, ?, ?, ?, ?)
        ''', (
            datetime.now().isoformat(),
            data_dict.get('batiment', ''),
            data_dict.get('redacteur', ''),
            data_dict.get('dateEssai', ''),
            json.dumps(data_dict),
            photo1_filename,
            photo2_filename
        ))
        conn.commit()
        report_id = c.lastrowid
        conn.close()
        
        return jsonify({'success': True, 'id': report_id, 'message': 'Rapport enregistré avec succès !'})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        username = request.form.get('username')
        password = request.form.get('password')
        if username == ADMIN_USERNAME and check_password_hash(ADMIN_PASSWORD_HASH, password):
            user = User('1')
            login_user(user)
            return redirect(url_for('admin'))
        return render_template('login.html', error='Identifiants incorrects')
    return render_template('login.html')

@app.route('/logout')
@login_required
def logout():
    logout_user()
    return redirect(url_for('login'))

@app.route('/admin')
@login_required
def admin():
    conn = get_db()
    c = conn.cursor()
    c.execute('SELECT id, submitted_at, batiment, redacteur, date_essai FROM reports ORDER BY submitted_at DESC')
    reports = c.fetchall()
    conn.close()
    return render_template('admin.html', reports=reports)

@app.route('/admin/view/<int:report_id>')
@login_required
def view_report(report_id):
    conn = get_db()
    c = conn.cursor()
    c.execute('SELECT * FROM reports WHERE id = ?', (report_id,))
    report = c.fetchone()
    conn.close()
    if not report:
        return 'Rapport non trouvé', 404
    data = json.loads(report['data_json'])
    return render_template('view.html', report=report, data=data)

@app.route('/admin/delete/<int:report_id>', methods=['POST'])
@login_required
def delete_report(report_id):
    conn = get_db()
    c = conn.cursor()
    # Optionnel : supprimer aussi les photos associées
    c.execute('SELECT photo1, photo2 FROM reports WHERE id = ?', (report_id,))
    row = c.fetchone()
    if row:
        for photo in (row['photo1'], row['photo2']):
            if photo:
                path = os.path.join(app.config['UPLOAD_FOLDER'], photo)
                if os.path.exists(path):
                    os.remove(path)
    c.execute('DELETE FROM reports WHERE id = ?', (report_id,))
    conn.commit()
    conn.close()
    return redirect(url_for('admin'))


@app.route('/admin/export')
@login_required
def export_excel():
    ids = request.args.getlist('ids')
    conn = get_db()
    c = conn.cursor()
    if ids:
        placeholders = ','.join('?' * len(ids))
        c.execute(f'SELECT * FROM reports WHERE id IN ({placeholders}) ORDER BY submitted_at DESC', ids)
    else:
        c.execute('SELECT * FROM reports ORDER BY submitted_at DESC')
    reports = c.fetchall()
    conn.close()

    if not reports:
        return "Aucun rapport sélectionné", 400

    # Charger le modèle
    template_path = os.path.join(os.path.dirname(__file__), 'template.xlsx')
    wb = load_workbook(template_path)
    template_ws = wb.active

    # Pour plusieurs rapports, on duplique la feuille
    for idx, r in enumerate(reports):
        data = json.loads(r['data_json'])
        if idx == 0:
            ws = template_ws
            ws.title = f"Rapport_{r['id']}"
        else:
            ws = wb.copy_worksheet(template_ws)
            ws.title = f"Rapport_{r['id']}"

        # Défusionner pour pouvoir écrire
        merged_ranges = list(ws.merged_cells.ranges)
        for merged in merged_ranges:
            ws.unmerge_cells(str(merged))

        # === EN-TÊTE ===
        ws['B10'] = f"COMPTE RENDU DES ESSAIS GE \nBâtiment : {data.get('batiment', '')}\nLe : {data.get('dateEssai', '')}"
        ws['F6'] = f"Rédacteur : {data.get('redacteur', '')}"

        # === RÉSUMÉ ===
        ws['B19'] = f"La consommation totale lors de ces essais a été de {data.get('consoTotale', 0)} litres."

        # === PROBLÈMES ===
        ws['B23'] = data.get('problemes', '')

        # === CUVES DÉBUT ===
        ws['C53'] = data.get('cuve1_debut', '')
        ws['C54'] = data.get('cuve2_debut', '')

        # === COMPTEURS HORAIRES DÉBUT ===
        ws['C62'] = data.get('gea_h_debut', '')
        ws['D62'] = data.get('gea_l_debut', '')
        ws['C63'] = data.get('geb_h_debut', '')
        ws['D63'] = data.get('geb_l_debut', '')
        ws['C64'] = data.get('gec_h_debut', '')
        ws['D64'] = data.get('gec_l_debut', '')
        ws['C65'] = data.get('ged_h_debut', '')
        ws['D65'] = data.get('ged_l_debut', '')

        # === COMPTEURS ÉLECTRIQUES DÉBUT ===
        ws['C72'] = data.get('gea_kwh_debut', '')
        ws['E72'] = data.get('gea_kvah_debut', '')
        ws['C73'] = data.get('geb_kwh_debut', '')
        ws['E73'] = data.get('geb_kvah_debut', '')
        ws['C74'] = data.get('gec_kwh_debut', '')
        ws['E74'] = data.get('gec_kvah_debut', '')
        ws['C75'] = data.get('ged_kwh_debut', '')
        ws['E75'] = data.get('ged_kvah_debut', '')

        # === GE A ===
        ws['C84'] = data.get('gea_h_run', '')
        ws['D84'] = data.get('gea_l_run', '')
        ws['D87'] = data.get('gea_v1', '')
        ws['E87'] = data.get('gea_v2', '')
        ws['F87'] = data.get('gea_v3', '')
        ws['D88'] = data.get('gea_i1', '')
        ws['E88'] = data.get('gea_i2', '')
        ws['F88'] = data.get('gea_i3', '')
        ws['D89'] = data.get('gea_freq', '')
        ws['D90'] = data.get('gea_cos', '')
        ws['D91'] = data.get('gea_kw', '')
        ws['D92'] = data.get('gea_kva', '')
        ws['D93'] = data.get('gea_kvar', '')
        ws['D94'] = data.get('gea_kwh_run', '')
        ws['D95'] = data.get('gea_kvah_run', '')
        ws['B100'] = data.get('gea_turbo', '')
        ws['C100'] = data.get('gea_huile', '')
        ws['D100'] = data.get('gea_fioul', '')
        ws['E100'] = data.get('gea_eau', '')
        ws['F100'] = data.get('gea_ech', '')

        # === GE B ===
        ws['C107'] = data.get('geb_h_run', '')
        ws['D107'] = data.get('geb_l_run', '')
        ws['D110'] = data.get('geb_v1', '')
        ws['E110'] = data.get('geb_v2', '')
        ws['F110'] = data.get('geb_v3', '')
        ws['D111'] = data.get('geb_i1', '')
        ws['E111'] = data.get('geb_i2', '')
        ws['F111'] = data.get('geb_i3', '')
        ws['D112'] = data.get('geb_freq', '')
        ws['D113'] = data.get('geb_cos', '')
        ws['D114'] = data.get('geb_kw', '')
        ws['D115'] = data.get('geb_kva', '')
        ws['D116'] = data.get('geb_kvar', '')
        ws['D117'] = data.get('geb_kwh_run', '')
        ws['D118'] = data.get('geb_kvah_run', '')
        ws['B123'] = data.get('geb_turbo', '')
        ws['C123'] = data.get('geb_huile', '')
        ws['D123'] = data.get('geb_fioul', '')
        ws['E123'] = data.get('geb_eau', '')
        ws['F123'] = data.get('geb_ech', '')

        # === GE C ===
        ws['C131'] = data.get('gec_h_run', '')
        ws['D131'] = data.get('gec_l_run', '')
        ws['D134'] = data.get('gec_v1', '')
        ws['E134'] = data.get('gec_v2', '')
        ws['F134'] = data.get('gec_v3', '')
        ws['D135'] = data.get('gec_i1', '')
        ws['E135'] = data.get('gec_i2', '')
        ws['F135'] = data.get('gec_i3', '')
        ws['D136'] = data.get('gec_freq', '')
        ws['D137'] = data.get('gec_cos', '')
        ws['D138'] = data.get('gec_kw', '')
        ws['D139'] = data.get('gec_kva', '')
        ws['D140'] = data.get('gec_kvar', '')
        ws['D141'] = data.get('gec_kwh_run', '')
        ws['D142'] = data.get('gec_kvah_run', '')
        ws['B147'] = data.get('gec_turbo', '')
        ws['C147'] = data.get('gec_huile', '')
        ws['D147'] = data.get('gec_fioul', '')
        ws['E147'] = data.get('gec_eau', '')
        ws['F147'] = data.get('gec_ech', '')

        # === GE D ===
        ws['C154'] = data.get('ged_h_run', '')
        ws['D154'] = data.get('ged_l_run', '')
        ws['D157'] = data.get('ged_v1', '')
        ws['E157'] = data.get('ged_v2', '')
        ws['F157'] = data.get('ged_v3', '')
        ws['D158'] = data.get('ged_i1', '')
        ws['E158'] = data.get('ged_i2', '')
        ws['F158'] = data.get('ged_i3', '')
        ws['D159'] = data.get('ged_freq', '')
        ws['D160'] = data.get('ged_cos', '')
        ws['D161'] = data.get('ged_kw', '')
        ws['D162'] = data.get('ged_kva', '')
        ws['D163'] = data.get('ged_kvar', '')
        ws['D164'] = data.get('ged_kwh_run', '')
        ws['D165'] = data.get('ged_kvah_run', '')
        ws['B170'] = data.get('ged_turbo', '')
        ws['C170'] = data.get('ged_huile', '')
        ws['D170'] = data.get('ged_fioul', '')
        ws['E170'] = data.get('ged_eau', '')
        ws['F170'] = data.get('ged_ech', '')

        # === FIN DES OPÉRATIONS ===
        ws['C179'] = data.get('cuve1_fin', '')
        ws['C180'] = data.get('cuve2_fin', '')

        # Compteurs horaires fin
        ws['C187'] = data.get('gea_h_fin', '')
        ws['D187'] = data.get('gea_l_fin', '')
        ws['C188'] = data.get('geb_h_fin', '')
        ws['D188'] = data.get('geb_l_fin', '')
        ws['C189'] = data.get('gec_h_fin', '')
        ws['D189'] = data.get('gec_l_fin', '')
        ws['C190'] = data.get('ged_h_fin', '')
        ws['D190'] = data.get('ged_l_fin', '')

        # Compteurs électriques fin
        ws['C196'] = data.get('gea_kwh_fin', '')
        ws['D196'] = data.get('gea_kvah_fin', '')
        ws['C197'] = data.get('geb_kwh_fin', '')
        ws['D197'] = data.get('geb_kvah_fin', '')
        ws['C198'] = data.get('gec_kwh_fin', '')
        ws['D198'] = data.get('gec_kvah_fin', '')
        ws['C199'] = data.get('ged_kwh_fin', '')
        ws['D199'] = data.get('ged_kvah_fin', '')

        # Consommations
        ws['C204'] = data.get('gea_conso', '')
        ws['C205'] = data.get('geb_conso', '')
        ws['C206'] = data.get('gec_conso', '')
        ws['C207'] = data.get('ged_conso', '')
        ws['B210'] = f"Soit une consommation de {data.get('consoTotale', 0)} litres pendant les essais."

        # Commentaires
        ws['B213'] = data.get('commentaires', '')

        # === FUSIONS DEMANDÉES ===
        # B87:B95 avec C
        for row in range(87, 96):
            ws.merge_cells(f'B{row}:C{row}')

        # B110:B118 avec C
        for row in range(110, 119):
            ws.merge_cells(f'B{row}:C{row}')

        # B134:B142 avec C
        for row in range(134, 143):
            ws.merge_cells(f'B{row}:C{row}')

        # B157:B165 avec C
        for row in range(157, 166):
            ws.merge_cells(f'B{row}:C{row}')

        # B202 avec C
        ws.merge_cells('B202:C202')

        # B210 jusqu'à E210
        ws.merge_cells('B210:E210')

        # === FUSIONS SUPPLÉMENTAIRES ===
        ws.merge_cells('B18:G18')
        ws.merge_cells('B22:G22')

        # === SUPPRESSION DE LIGNES ===
        for row_num in sorted([216, 208, 200, 191], reverse=True):
            ws.delete_rows(row_num)
        # Supprimer les lignes 24 à 30 (après les autres suppressions, ajuster si besoin)
        ws.delete_rows(24, 7)  # Supprime 7 lignes à partir de la ligne 24

        # === CENTRAGE DE TOUT LE TEXTE ===
        center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=7):
            for cell in row:
                if cell.value is not None:
                    cell.alignment = center_align

        # === ENCADREMENTS (bordures extérieures uniquement) ===
        thin = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        # Bordure extérieure B10:G14
        for row in range(10, 15):
            for col in range(2, 8):  # B à G
                cell = ws.cell(row=row, column=col)
                border = Border()
                if row == 10:
                    border.top = Side(style='thin')
                if row == 14:
                    border.bottom = Side(style='thin')
                if col == 2:
                    border.left = Side(style='thin')
                if col == 7:
                    border.right = Side(style='thin')
                cell.border = border

        # Bordure extérieure B16:G20
        for row in range(16, 21):
            for col in range(2, 8):
                cell = ws.cell(row=row, column=col)
                border = Border()
                if row == 16:
                    border.top = Side(style='thin')
                if row == 20:
                    border.bottom = Side(style='thin')
                if col == 2:
                    border.left = Side(style='thin')
                if col == 7:
                    border.right = Side(style='thin')
                cell.border = border

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"rapports_ge_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return send_file(output, download_name=filename, as_attachment=True,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.route('/uploads/<filename>')
@login_required
def uploaded_file(filename):
    return send_from_directory(app.config['UPLOAD_FOLDER'], filename)

if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=5000)
